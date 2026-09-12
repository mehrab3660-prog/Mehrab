"""
برنامه‌ی گرافیکی استخراج پلاک و وضعیت مخزن از سامانه سیمفا (gas.symfa.ir)

نصب یک‌بار (خط فرمان):
    pip install requests beautifulsoup4 openpyxl

اجرا:
    فایل run.vbs (یا run.bat) رو دابل‌کلیک کن.

نام‌کاربری و رمز عبورت رو توی برنامه وارد کن؛ برنامه خودش با همون‌ها
لاگین می‌کنه و دیگه نیازی به کپی‌کردن کوکی از DevTools نیست.
"""

import os
import re
import threading
import tkinter as tk
from datetime import datetime, timedelta
from tkinter import ttk, scrolledtext, messagebox
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = "https://gas.symfa.ir"
LOGIN_URL = f"{BASE}/TestCenters/Home/Login"
LIST_URL = f"{BASE}/TestCenters/GasReception"
PRINT_URL = f"{BASE}/TestCenters/GasReception/PrintResult?ReceptionId={{}}"

APP_DIR = os.path.dirname(os.path.abspath(__file__))
CREDS_FILE = os.path.join(APP_DIR, "last_login.txt")
OUTPUT_FILE = os.path.join(APP_DIR, "result.xlsx")

PERSIAN_WEEKDAYS = ["دوشنبه", "سه‌شنبه", "چهارشنبه", "پنجشنبه", "جمعه", "شنبه", "یکشنبه"]
PERSIAN_MONTHS = [
    "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
    "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند",
]


def gregorian_to_jalali(gy, gm, gd):
    """تبدیل تاریخ میلادی به شمسی (بدون نیاز به کتابخونه‌ی جداگانه)."""
    g_d_m = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    if gy > 1600:
        jy = 979
        gy -= 1600
    else:
        jy = 0
        gy -= 621
    gy2 = gy + 1 if gm > 2 else gy
    days = (
        365 * gy
        + (gy2 + 3) // 4
        - (gy2 + 99) // 100
        + (gy2 + 399) // 400
        - 80
        + gd
        + g_d_m[gm - 1]
    )
    jy += 33 * (days // 12053)
    days %= 12053
    jy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jy += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jm = 1 + days // 31
        jd = 1 + (days % 31)
    else:
        jm = 7 + (days - 186) // 30
        jd = 1 + ((days - 186) % 30)
    return jy, jm, jd


def jalali_to_gregorian(jy, jm, jd):
    """تبدیل تاریخ شمسی به میلادی (معکوس gregorian_to_jalali)."""
    jy += 1595
    days = -355668 + (365 * jy) + (jy // 33) * 8 + (((jy % 33) + 3) // 4) + jd
    days += (jm - 1) * 31 if jm < 7 else ((jm - 7) * 30) + 186
    gy = 400 * (days // 146097)
    days %= 146097
    if days > 36524:
        days -= 1
        gy += 100 * (days // 36524)
        days %= 36524
        days += 1
    gy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        gy += (days - 1) // 365
        days = (days - 1) % 365
    gd = days + 1
    leap = gy % 4 == 0 and (gy % 100 != 0 or gy % 400 == 0)
    g_days_in_month = [31, 29 if leap else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    gm = 0
    while gm < 12 and gd > g_days_in_month[gm]:
        gd -= g_days_in_month[gm]
        gm += 1
    return gy, gm + 1, gd


def jalali_weekday_label(jy, jm, jd):
    """اسم روز هفته + تاریخ شمسی رو برای یه تاریخ شمسیِ دلخواه می‌سازه."""
    gy, gm, gd = jalali_to_gregorian(jy, jm, jd)
    weekday = PERSIAN_WEEKDAYS[datetime(gy, gm, gd).weekday()]
    return f"{weekday}، {jd} {PERSIAN_MONTHS[jm - 1]} {jy}"


def today_jalali_string():
    now = datetime.now()
    jy, jm, jd = gregorian_to_jalali(now.year, now.month, now.day)
    return jalali_weekday_label(jy, jm, jd)


def jalali_date_str(dt):
    """تاریخ رو به همون فرمتی که فیلتر سایت می‌خواد برمی‌گردونه: 1405/06/19"""
    jy, jm, jd = gregorian_to_jalali(dt.year, dt.month, dt.day)
    return f"{jy}/{jm:02d}/{jd:02d}"


def report_date_label(from_date_str, to_date_str):
    """برچسب تاریخ گزارش رو از روی بازه‌ی انتخاب‌شده (نه تاریخ اجرای برنامه) می‌سازه."""
    try:
        fy, fm, fd = (int(x) for x in from_date_str.split("/"))
        ty, tm, td = (int(x) for x in to_date_str.split("/"))
    except ValueError:
        return f"{from_date_str} تا {to_date_str}"

    if (fy, fm, fd) == (ty, tm, td):
        return jalali_weekday_label(fy, fm, fd)
    return f"از {jalali_weekday_label(fy, fm, fd)} تا {jalali_weekday_label(ty, tm, td)}"


def jalali_days_ago(n):
    return jalali_date_str(datetime.now() - timedelta(days=n))

# پالت رنگی - تم تیره‌ی طلایی/سرمه‌ای
COLOR_BG = "#0e1117"
COLOR_HEADER_FROM = "#151a27"
COLOR_HEADER_TO = "#241a3a"
COLOR_HEADER_TEXT = "#e8c876"
COLOR_ACCENT = "#c9a227"
COLOR_ACCENT_HOVER = "#e0b830"
COLOR_ACCENT_TEXT = "#171308"
COLOR_CARD = "#171b26"
COLOR_BORDER = "#2a3040"
COLOR_TEXT = "#e8e6e3"
COLOR_MUTED = "#8b94a7"
COLOR_FIELD_BG = "#0f1320"
FONT_FAMILY = "Segoe UI"


def draw_car_icon(canvas, cx, cy, scale, color, bg="#0e1117"):
    """یه سیلوئت ماشین ساده و مسطح (flat) روی Canvas می‌کشه."""
    # بدنه‌ی اصلی
    canvas.create_rectangle(
        cx - 24 * scale, cy - 2 * scale, cx + 24 * scale, cy + 4 * scale,
        fill=color, outline=color,
    )
    # کابین (سقف)
    canvas.create_polygon(
        cx - 14 * scale, cy - 2 * scale,
        cx - 9 * scale, cy - 10 * scale,
        cx + 6 * scale, cy - 10 * scale,
        cx + 10 * scale, cy - 2 * scale,
        fill=color, outline=color,
    )
    # شیشه‌ی جلو/عقب
    canvas.create_line(
        cx - 10 * scale, cy - 9 * scale, cx - 6 * scale, cy - 2.5 * scale,
        fill=bg, width=max(1, int(1.2 * scale)),
    )
    canvas.create_line(
        cx + 2 * scale, cy - 2.5 * scale, cx + 5 * scale, cy - 9 * scale,
        fill=bg, width=max(1, int(1.2 * scale)),
    )
    # چراغ جلو
    canvas.create_oval(
        cx + 20 * scale, cy - 1 * scale, cx + 24 * scale, cy + 1.5 * scale,
        fill="#fff3d0", outline="",
    )
    # چرخ‌ها (توپر، بدون نقطه‌ی داخلی)
    r = 4 * scale
    wy = cy + 4 * scale
    for wx in (-13, 13):
        wcx = cx + wx * scale
        canvas.create_oval(wcx - r, wy - r, wcx + r, wy + r, fill="#3a4155", outline=bg, width=max(1, int(scale)))


def draw_gradient(canvas, width, height, color_from, color_to):
    """یه گرادیان افقی روی یه Canvas می‌کشه (تک‌رنگ ttk نداره)."""

    def hex_to_rgb(h):
        h = h.lstrip("#")
        return tuple(int(h[i : i + 2], 16) for i in (0, 2, 4))

    r1, g1, b1 = hex_to_rgb(color_from)
    r2, g2, b2 = hex_to_rgb(color_to)
    steps = max(width, 1)
    for i in range(steps):
        t = i / steps
        r = int(r1 + (r2 - r1) * t)
        g = int(g1 + (g2 - g1) * t)
        b = int(b1 + (b2 - b1) * t)
        color = f"#{r:02x}{g:02x}{b:02x}"
        canvas.create_line(i, 0, i, height, fill=color)


def login(username, password):
    """با نام‌کاربری/رمز وارد می‌شود و یک session با کوکی معتبر برمی‌گرداند."""
    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0"})

    # یه بار صفحه رو باز می‌کنیم تا کوکی اولیه‌ی سشن گرفته بشه
    session.get(BASE, timeout=30)

    resp = session.post(
        LOGIN_URL,
        data={"UserName": username, "Password": password, "IsSoft": ""},
        timeout=30,
        allow_redirects=True,
    )
    resp.raise_for_status()
    return session


def fetch_reception_list(session, start_date, end_date):
    """لیست پذیرش‌ها رو با فیلتر بازه‌ی تاریخ (شمسی، فرمت 1405/06/19) می‌گیره."""
    resp = session.post(
        LIST_URL,
        data={
            "StartDate": start_date,
            "EndDate": end_date,
            "ReceptionId": "",
            "VIN": "",
        },
        timeout=30,
    )
    resp.raise_for_status()
    return resp


def extract_reception_rows(html):
    """از جدول صفحه‌ی لیست، (پلاک، کد پذیرش) هر ردیف را برمی‌گرداند."""
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table")
    if table is None or table.find("thead") is None or table.find("tbody") is None:
        raise RuntimeError(
            "جدول پذیرش‌ها پیدا نشد - احتمالاً نام‌کاربری یا رمز اشتباهه."
        )

    headers = [th.get_text(strip=True) for th in table.find("thead").find_all("th")]
    plate_idx = next((i for i, h in enumerate(headers) if "پلاک" in h), None)
    code_idx = next((i for i, h in enumerate(headers) if "کد پذیرش" in h), None)
    if plate_idx is None or code_idx is None:
        raise RuntimeError(f"ستون‌های پلاک/کد پذیرش پیدا نشد. هدرها: {headers}")

    rows = []
    for tr in table.find("tbody").find_all("tr"):
        cells = [td.get_text(strip=True) for td in tr.find_all("td")]
        if len(cells) <= max(plate_idx, code_idx):
            continue
        plate = cells[plate_idx]
        code = cells[code_idx]
        if code.isdigit():
            rows.append((plate, code))
    return rows


def extract_pagination_urls(html, current_url):
    """آدرس کامل هر صفحه‌ی صفحه‌بندی رو برمی‌گردونه (شماره صفحه -> URL کامل)،
    تا فیلترهای اعمال‌شده (مثل بازه‌ی تاریخ) موقع رفتن به صفحه‌ی بعد حفظ بشن."""
    soup = BeautifulSoup(html, "html.parser")
    urls = {1: current_url}
    for a in soup.find_all("a", href=True):
        m = re.search(r"[?&]Page=(\d+)", a["href"])
        if m:
            urls[int(m.group(1))] = urljoin(current_url, a["href"])
    return urls


def extract_tanks(html):
    """وضعیت هر مخزن را استخراج می‌کند (یک ورودی به ازای هر مخزن)."""
    soup = BeautifulSoup(html, "html.parser")
    tanks = []
    for span in soup.find_all("span"):
        text = span.get_text(strip=True)
        if "مخزن" not in text:
            continue
        if "تایید" in text:
            tanks.append("تایید")
        elif "مردود" in text:
            tanks.append("مردود")
    return tanks


def extract_plate(html):
    soup = BeautifulSoup(html, "html.parser")
    cell = soup.find("td", class_="PlaqueNumber")
    return cell.get_text(strip=True) if cell else ""


# رنگ‌های خروجی اکسل (هم‌رنگ با تم برنامه)
XLS_NAVY = "151A27"
XLS_LIGHT = "F4F1E6"
XLS_GREEN_FILL = "E5F5E6"
XLS_GREEN_FONT = "1B7A2E"
XLS_RED_FILL = "FBE7E7"
XLS_RED_FONT = "B3241C"


def add_report_block(ws, start_col, start_row, title, headers, data_rows):
    """یه بلوک (عنوان + هدر + ردیف‌ها) رو کنار بلوک‌های قبلی (ستون به ستون) اضافه می‌کنه
    و شماره‌ی ستون شروع بلوک بعدی رو برمی‌گردونه."""
    n_cols = len(headers)
    end_col = start_col + n_cols - 1
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)
    tcell = ws.cell(row=start_row, column=start_col, value=f"{title} ({len(data_rows)})")
    tcell.font = Font(bold=True, size=14, color="FFFFFF")
    tcell.fill = PatternFill("solid", fgColor=XLS_NAVY)
    tcell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[start_row].height = 22

    header_row = start_row + 1
    for i, h in enumerate(headers):
        c = ws.cell(row=header_row, column=start_col + i, value=h)
        c.font = Font(bold=True, color=XLS_NAVY, size=12)
        c.fill = PatternFill("solid", fgColor=XLS_LIGHT)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    ws.row_dimensions[header_row].height = 20

    for offset, row in enumerate(data_rows):
        r_i = header_row + 1 + offset
        ws.row_dimensions[r_i].height = 18
        for i, val in enumerate(row):
            cell = ws.cell(row=r_i, column=start_col + i, value=val)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(size=12)
            if val == "تایید":
                cell.font = Font(bold=True, size=12, color=XLS_GREEN_FONT)
                cell.fill = PatternFill("solid", fgColor=XLS_GREEN_FILL)
            elif val == "مردود":
                cell.font = Font(bold=True, size=12, color=XLS_RED_FONT)
                cell.fill = PatternFill("solid", fgColor=XLS_RED_FILL)

    return end_col + 2  # یک ستون خالی به‌عنوان فاصله


def build_report_workbook(
    one_tank_rows, two_tank_rows, rejected_one_tank_rows, rejected_two_tank_rows,
    other_rows, date_label, stack_threshold=10,
):
    """یه فایل اکسل تک‌شیت با بلوک‌های تفکیک‌شده (تک‌مخزن/دومخزن/مردودی تک‌مخزن/
    مردودی جفت‌مخزن/سایر) کنار هم، همه روی یک صفحه‌ی A5. هر پلاک فقط توی یکی از
    بلوک‌ها میاد.

    اگه تعداد ردیف‌های یه بلوک (تک‌مخزن یا دومخزن) کم باشه (<= stack_threshold)،
    مردودیِ همون دسته به‌جای بلوک جدا و کنار هم، زیر همون بلوک (توی همون
    ستون‌ها) میاد.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "گزارش"
    ws.sheet_view.rightToLeft = True

    three_col_headers = ["ردیف", "پلاک", "وضعیت"]

    date_row = 1
    block_row = date_row + 1
    col = 1

    def place_block(title, headers, rows, start_row=block_row, start_col=None, advance_col=True):
        nonlocal col
        scol = col if start_col is None else start_col
        next_col = add_report_block(ws, scol, start_row, title, headers, rows)
        for c in range(scol, next_col - 1):
            width = 7 if headers[c - scol] == "ردیف" else 15
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.column_dimensions[get_column_letter(next_col - 1)].width = 2
        end_row = start_row + 1 + len(rows)  # عنوان + هدر + ردیف‌ها
        if advance_col:
            col = next_col
        return end_row, scol

    def place_with_rejected(title, rows, rejected_title, rejected_rows):
        stack = bool(rows) and bool(rejected_rows) and len(rows) <= stack_threshold
        if rows:
            start_col = col
            end_row, _ = place_block(title, three_col_headers, rows)
            if stack:
                place_block(
                    rejected_title, three_col_headers, rejected_rows,
                    start_row=end_row + 2, start_col=start_col, advance_col=False,
                )
        if rejected_rows and not stack:
            place_block(rejected_title, three_col_headers, rejected_rows)

    place_with_rejected("تک‌مخزن", one_tank_rows, "مردودی تک‌مخزن", rejected_one_tank_rows)
    place_with_rejected("دو‌مخزن", two_tank_rows, "مردودی جفت‌مخزن", rejected_two_tank_rows)

    if other_rows:
        place_block("سایر", ["ردیف", "پلاک", "توضیح"], other_rows)

    total_cols = max(col - 2, 1)
    ws.merge_cells(start_row=date_row, start_column=1, end_row=date_row, end_column=total_cols)
    date_cell = ws.cell(row=date_row, column=1, value=date_label)
    date_cell.font = Font(bold=True, size=12, color=XLS_NAVY)
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[date_row].height = 22

    ws.page_setup.paperSize = ws.PAPERSIZE_A5
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3

    return wb


class App:
    def __init__(self, root):
        self.root = root
        root.title("استخراج پلاک")
        root.geometry("700x600")
        root.minsize(560, 460)
        root.configure(bg=COLOR_BG)
        self._set_window_icon(root)

        style = ttk.Style(root)
        style.theme_use("clam")
        style.configure("Card.TFrame", background=COLOR_CARD)
        style.configure(
            "Field.TLabel",
            background=COLOR_CARD,
            foreground=COLOR_ACCENT,
            font=(FONT_FAMILY, 10, "bold"),
        )
        style.configure(
            "Status.TLabel",
            background=COLOR_CARD,
            foreground=COLOR_MUTED,
            font=(FONT_FAMILY, 9),
        )
        style.configure(
            "Accent.TButton",
            background=COLOR_ACCENT,
            foreground=COLOR_ACCENT_TEXT,
            font=(FONT_FAMILY, 11, "bold"),
            padding=(18, 10),
            borderwidth=0,
        )
        style.map(
            "Accent.TButton",
            background=[("active", COLOR_ACCENT_HOVER), ("disabled", "#4a4a42")],
            foreground=[("disabled", "#8a8a80")],
        )
        style.configure(
            "Field.TEntry",
            fieldbackground=COLOR_FIELD_BG,
            foreground=COLOR_TEXT,
            insertcolor=COLOR_TEXT,
            bordercolor=COLOR_BORDER,
            lightcolor=COLOR_BORDER,
            darkcolor=COLOR_BORDER,
            padding=8,
        )
        style.map("Field.TEntry", bordercolor=[("focus", COLOR_ACCENT)])
        style.configure(
            "Accent.Horizontal.TProgressbar",
            troughcolor=COLOR_FIELD_BG,
            background=COLOR_ACCENT,
            thickness=6,
        )
        style.configure(
            "Preset.TButton",
            background=COLOR_FIELD_BG,
            foreground=COLOR_ACCENT,
            font=(FONT_FAMILY, 9),
            padding=(10, 5),
            borderwidth=1,
        )
        style.map("Preset.TButton", background=[("active", COLOR_BORDER)])

        # ---------- هدر (گرادیان سرمه‌ای/بنفش با عنوان طلایی) ----------
        header_h = 92
        logo_size = 60
        header = tk.Canvas(root, height=header_h, highlightthickness=0, bd=0)
        header.pack(fill="x")

        def render_header(event=None):
            header.delete("all")
            w = header.winfo_width() or root.winfo_width() or 700
            draw_gradient(header, w, header_h, COLOR_HEADER_FROM, COLOR_HEADER_TO)
            header.create_line(0, header_h - 1, w, header_h - 1, fill=COLOR_ACCENT, width=2)
            logo_cx = w - 20 - logo_size // 2
            logo_cy = header_h // 2
            draw_car_icon(header, logo_cx, logo_cy, 1.6, COLOR_ACCENT, bg=COLOR_HEADER_TO)
            text_right = logo_cx - 44
            header.create_text(
                text_right,
                header_h // 2,
                text="استخراج پلاک",
                fill=COLOR_HEADER_TEXT,
                font=(FONT_FAMILY, 20, "bold"),
                anchor="e",
            )

        header.bind("<Configure>", render_header)

        # ---------- بدنه ----------
        body = tk.Frame(root, bg=COLOR_BG)
        body.pack(fill="both", expand=True, padx=18, pady=16)

        card = tk.Frame(body, bg=COLOR_CARD, highlightbackground=COLOR_BORDER, highlightthickness=1)
        card.pack(fill="both", expand=True)

        inner = tk.Frame(card, bg=COLOR_CARD)
        inner.pack(fill="both", expand=True, padx=18, pady=16)

        saved_user, saved_pass = self._load_creds()

        login_row = tk.Frame(inner, bg=COLOR_CARD)
        login_row.pack(fill="x", pady=(0, 4))

        user_col = tk.Frame(login_row, bg=COLOR_CARD)
        user_col.pack(side="right", fill="x", expand=True, padx=(8, 0))
        ttk.Label(user_col, text="نام‌کاربری", style="Field.TLabel").pack(anchor="e")
        self.user_entry = ttk.Entry(user_col, style="Field.TEntry", justify="right")
        self.user_entry.pack(fill="x", pady=(4, 0))
        self.user_entry.insert(0, saved_user)

        pass_col = tk.Frame(login_row, bg=COLOR_CARD)
        pass_col.pack(side="right", fill="x", expand=True)
        ttk.Label(pass_col, text="رمز عبور", style="Field.TLabel").pack(anchor="e")
        self.pass_entry = ttk.Entry(pass_col, style="Field.TEntry", justify="right", show="•")
        self.pass_entry.pack(fill="x", pady=(4, 0))
        self.pass_entry.insert(0, saved_pass)

        date_row = tk.Frame(inner, bg=COLOR_CARD)
        date_row.pack(fill="x", pady=(10, 4))

        today_str = jalali_date_str(datetime.now())

        to_col = tk.Frame(date_row, bg=COLOR_CARD)
        to_col.pack(side="right", fill="x", expand=True, padx=(8, 0))
        ttk.Label(to_col, text="تا تاریخ", style="Field.TLabel").pack(anchor="e")
        self.to_date_entry = ttk.Entry(to_col, style="Field.TEntry", justify="right")
        self.to_date_entry.pack(fill="x", pady=(4, 0))
        self.to_date_entry.insert(0, today_str)

        from_col = tk.Frame(date_row, bg=COLOR_CARD)
        from_col.pack(side="right", fill="x", expand=True)
        ttk.Label(from_col, text="از تاریخ", style="Field.TLabel").pack(anchor="e")
        self.from_date_entry = ttk.Entry(from_col, style="Field.TEntry", justify="right")
        self.from_date_entry.pack(fill="x", pady=(4, 0))
        self.from_date_entry.insert(0, today_str)

        preset_row = tk.Frame(inner, bg=COLOR_CARD)
        preset_row.pack(fill="x", pady=(6, 0))
        ttk.Button(
            preset_row, text="امروز", style="Preset.TButton",
            command=lambda: self._set_date_range(0, 0),
        ).pack(side="right", padx=(6, 0))
        ttk.Button(
            preset_row, text="دیروز", style="Preset.TButton",
            command=lambda: self._set_date_range(1, 1),
        ).pack(side="right", padx=(6, 0))
        ttk.Button(
            preset_row, text="۷ روز اخیر", style="Preset.TButton",
            command=lambda: self._set_date_range(7, 0),
        ).pack(side="right")

        action_row = tk.Frame(inner, bg=COLOR_CARD)
        action_row.pack(fill="x", pady=(14, 10))
        self.start_btn = ttk.Button(
            action_row, text="▶  شروع استخراج", style="Accent.TButton", command=self.start
        )
        self.start_btn.pack(side="right")

        self.status_var = tk.StringVar(value="آماده")
        ttk.Label(action_row, textvariable=self.status_var, style="Status.TLabel").pack(
            side="right", padx=(0, 14)
        )

        self.progress = ttk.Progressbar(
            inner, style="Accent.Horizontal.TProgressbar", mode="indeterminate"
        )
        self.progress.pack(fill="x", pady=(0, 12))

        ttk.Label(inner, text="گزارش اجرا", style="Field.TLabel").pack(anchor="e", fill="x")
        self.log_box = scrolledtext.ScrolledText(
            inner,
            font=("Consolas", 9),
            bg="#0f1720",
            fg="#d7ecf3",
            insertbackground="#d7ecf3",
            relief="flat",
            state="disabled",
        )
        self.log_box.pack(fill="both", expand=True, pady=(6, 0))

    def _set_date_range(self, from_days_ago, to_days_ago):
        self.from_date_entry.delete(0, "end")
        self.from_date_entry.insert(0, jalali_days_ago(from_days_ago))
        self.to_date_entry.delete(0, "end")
        self.to_date_entry.insert(0, jalali_days_ago(to_days_ago))

    def _set_window_icon(self, root):
        """آیکون پنجره/نوار وظیفه رو با یه لوگوی ساده‌ی ماشین می‌سازه (بدون فایل خارجی)."""
        try:
            size = 32
            img = tk.PhotoImage(width=size, height=size)
            img.put(COLOR_BG, to=(0, 0, size, size))
            img.put(COLOR_ACCENT, to=(10, 11, 22, 17))  # کابین
            img.put(COLOR_ACCENT, to=(4, 17, 28, 22))  # بدنه
            img.put("#3a4155", to=(7, 21, 13, 27))  # چرخ جلو
            img.put("#3a4155", to=(19, 21, 25, 27))  # چرخ عقب
            self._icon_img = img  # جلوگیری از garbage collection
            root.iconphoto(True, img)
        except Exception:
            pass

    def _load_creds(self):
        if os.path.exists(CREDS_FILE):
            with open(CREDS_FILE, "r", encoding="utf-8") as f:
                lines = f.read().splitlines()
            if len(lines) >= 2:
                return lines[0], lines[1]
        return "", ""

    def _save_creds(self, username, password):
        with open(CREDS_FILE, "w", encoding="utf-8") as f:
            f.write(username + "\n" + password)

    def log(self, text):
        self.log_box.config(state="normal")
        self.log_box.insert("end", text + "\n")
        self.log_box.see("end")
        self.log_box.config(state="disabled")

    def set_status(self, text):
        self.status_var.set(text)

    def start(self):
        username = self.user_entry.get().strip()
        password = self.pass_entry.get().strip()
        from_date = self.from_date_entry.get().strip()
        to_date = self.to_date_entry.get().strip()
        if not username or not password:
            messagebox.showerror("خطا", "نام‌کاربری و رمز عبور رو وارد کن.")
            return
        if not from_date or not to_date:
            messagebox.showerror("خطا", "از تاریخ و تا تاریخ رو وارد کن.")
            return
        self.start_btn.config(state="disabled")
        self.progress.start(12)
        self.set_status("در حال ورود...")
        threading.Thread(
            target=self.run, args=(username, password, from_date, to_date), daemon=True
        ).start()

    def run(self, username, password, from_date, to_date):
        try:
            self._save_creds(username, password)

            self.log("در حال ورود به حساب کاربری...")
            session = login(username, password)

            self.set_status("در حال گرفتن لیست پذیرش‌ها...")
            self.log(f"در حال گرفتن پذیرش‌های {from_date} تا {to_date}...")
            resp = fetch_reception_list(session, from_date, to_date)
            pagination = extract_pagination_urls(resp.text, resp.url)
            max_page = max(pagination)
            self.log(f"{max_page} صفحه پیدا شد.")

            rows = []
            seen_codes = set()
            for plate, code in extract_reception_rows(resp.text):
                if code not in seen_codes:
                    seen_codes.add(code)
                    rows.append((plate, code))

            for page in range(2, max_page + 1):
                self.set_status(f"در حال گرفتن صفحه‌ی {page} از {max_page}...")
                page_resp = session.get(pagination[page], timeout=30)
                page_resp.raise_for_status()
                for plate, code in extract_reception_rows(page_resp.text):
                    if code not in seen_codes:
                        seen_codes.add(code)
                        rows.append((plate, code))

            self.log(f"{len(rows)} پذیرش پیدا شد.")

            results = []
            for i, (plate, code) in enumerate(rows, start=1):
                self.set_status(f"در حال پردازش {i} از {len(rows)}...")
                try:
                    r = session.get(PRINT_URL.format(code), timeout=30)
                    r.raise_for_status()
                    actual_plate = extract_plate(r.text) or plate
                    tanks = extract_tanks(r.text)
                except Exception as item_err:
                    self.log(f"{code} - رد شد (خطا: {item_err})")
                    continue
                results.append(
                    {
                        "کد پذیرش": code,
                        "پلاک": actual_plate,
                        "تعداد مخزن": len(tanks),
                        "وضعیت مخزن ۱": tanks[0] if len(tanks) > 0 else "",
                        "وضعیت مخزن ۲": tanks[1] if len(tanks) > 1 else "",
                    }
                )
                self.log(f"{code} - {actual_plate} - {tanks}")

            rejected_codes = {
                r["کد پذیرش"] for r in results
                if "مردود" in (r["وضعیت مخزن ۱"], r["وضعیت مخزن ۲"])
            }

            # هر پلاک فقط توی یه دسته میاد - قبول/مردود و تک‌مخزن/دومخزن کاملاً تفکیک شده
            one_tank = [
                r for r in results
                if r["تعداد مخزن"] == 1 and r["کد پذیرش"] not in rejected_codes
            ]
            two_tank = [
                r for r in results
                if r["تعداد مخزن"] == 2 and r["کد پذیرش"] not in rejected_codes
            ]
            rejected_one_tank = [
                r for r in results
                if r["تعداد مخزن"] == 1 and r["کد پذیرش"] in rejected_codes
            ]
            rejected_two_tank = [
                r for r in results
                if r["تعداد مخزن"] == 2 and r["کد پذیرش"] in rejected_codes
            ]
            other = [r for r in results if r["تعداد مخزن"] not in (1, 2)]

            one_tank_rows = [
                [i, r["پلاک"], "تایید"] for i, r in enumerate(one_tank, start=1)
            ]
            two_tank_rows = [
                [i, r["پلاک"], "تایید"] for i, r in enumerate(two_tank, start=1)
            ]
            rejected_one_tank_rows = [
                [i, r["پلاک"], "مردود"] for i, r in enumerate(rejected_one_tank, start=1)
            ]
            rejected_two_tank_rows = [
                [i, r["پلاک"], "مردود"] for i, r in enumerate(rejected_two_tank, start=1)
            ]
            other_rows = [
                [i, r["پلاک"], f"{r['تعداد مخزن']} مخزن"]
                for i, r in enumerate(other, start=1)
            ]

            wb = build_report_workbook(
                one_tank_rows, two_tank_rows,
                rejected_one_tank_rows, rejected_two_tank_rows, other_rows,
                report_date_label(from_date, to_date),
            )
            wb.save(OUTPUT_FILE)
            self.log(f"ذخیره شد در: {OUTPUT_FILE}")
            self.log(
                f"تک‌مخزن: {len(one_tank)} - دومخزنه: {len(two_tank)} - "
                f"مردودی تک‌مخزن: {len(rejected_one_tank)} - مردودی جفت‌مخزن: {len(rejected_two_tank)}"
            )
            self.set_status(f"تمام شد - {len(results)} پذیرش")
            messagebox.showinfo(
                "تمام شد",
                f"{len(results)} پذیرش استخراج شد "
                f"({len(one_tank)} تک‌مخزن، {len(two_tank)} دومخزنه، "
                f"{len(rejected_one_tank)} مردودی تک‌مخزن، {len(rejected_two_tank)} مردودی جفت‌مخزن) "
                f"و در result.xlsx ذخیره شد.",
            )
            try:
                os.startfile(OUTPUT_FILE)
            except Exception:
                pass
        except Exception as e:
            self.log(f"خطا: {e}")
            self.set_status("خطا")
            messagebox.showerror("خطا", str(e))
        finally:
            self.progress.stop()
            self.start_btn.config(state="normal")


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
